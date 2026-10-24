import os
import uuid
from functools import wraps
from flask import Flask, render_template, request, redirect, url_for, flash, send_file
from flask_login import LoginManager, login_user, login_required, logout_user, current_user
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook, load_workbook
from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo  # 【新增】导入时区库

from db import init_db, get_db
from models import get_user_by_id, get_user_by_username
from services import child_service, class_service, user_service, record_service

app = Flask(__name__)
app.secret_key = "change-this-to-a-random-string"

app.config['UPLOAD_FOLDER'] = os.path.join(app.root_path, 'static', 'uploads')
app.config['MAX_CONTENT_LENGTH'] = 5 * 1024 * 1024

login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = "login"
login_manager.login_message = "请先登录"

@login_manager.user_loader
def load_user(user_id):
    return get_user_by_id(int(user_id))

def role_required(*allowed_roles):
    def decorator(f):
        @wraps(f)
        def wrapper(*args, **kwargs):
            if not current_user.is_authenticated:
                return redirect(url_for("login"))
            if current_user.role not in allowed_roles:
                flash("无权限访问该页面")
                return redirect(url_for("index"))
            return f(*args, **kwargs)
        return wrapper
    return decorator

# ---------- 认证 ----------
@app.route("/register", methods=["GET", "POST"])
def register():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        real_name = request.form.get("real_name", "")
        role = "parent"
        if not username or not password:
            flash("用户名和密码不能为空")
            return redirect(url_for("register"))
        if get_user_by_username(username):
            flash("用户名已存在")
            return redirect(url_for("register"))
        conn = get_db()
        conn.execute(
            "INSERT INTO users (username, password_hash, role, real_name) VALUES (?, ?, ?, ?)",
            (username, generate_password_hash(password), role, real_name),
        )
        conn.commit()
        conn.close()
        flash("注册成功，请登录。如需绑定孩子信息，请联系管理员。")
        return redirect(url_for("login"))
    return render_template("auth/register.html")

@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        user = get_user_by_username(username)
        if user and check_password_hash(user.password_hash, password):
            login_user(user)
            return redirect(url_for("index"))
        flash("用户名或密码错误")
    return render_template("auth/login.html")

@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))

# ---------- 首页 ----------
@app.route("/")
@login_required
def index():
    if current_user.role in ("admin", "teacher"):
        children = child_service.list_children()
        classes = class_service.list_classes()
        notices = get_notices()
        return render_template("index.html", child_count=len(children), class_count=len(classes), notices=notices)
    
    child = None
    photos = []
    if current_user.child_id:
        child = child_service.get_child(current_user.child_id)
        photos = get_photos_by_child(current_user.child_id)
    notices = get_notices()
    return render_template("index.html", child=child, notices=notices, photos=photos)

# ---------- 儿童管理 ----------
@app.route("/children")
@login_required
@role_required("admin", "teacher")
def children_list():
    keyword = request.args.get("q", "")
    class_id = request.args.get("class_id", type=int)
    children = child_service.list_children(keyword or None, class_id)
    classes = class_service.list_classes()
    return render_template("children/list.html", children=children, classes=classes,
                           keyword=keyword, selected_class=class_id)

@app.route("/children/add", methods=["GET", "POST"])
@login_required
@role_required("admin", "teacher")
def children_add():
    classes = class_service.list_classes()
    if request.method == "POST":
        child_service.add_child(
            request.form["name"], request.form["gender"], request.form["birthday"],
            request.form.get("class_id") or None, request.form["guardian"],
            request.form["phone"], request.form.get("allergy", ""),
        )
        flash("添加成功")
        return redirect(url_for("children_list"))
    return render_template("children/form.html", child=None, classes=classes)

@app.route("/children/<int:child_id>/edit", methods=["GET", "POST"])
@login_required
@role_required("admin", "teacher")
def children_edit(child_id):
    child = child_service.get_child(child_id)
    if not child:
        flash("儿童不存在")
        return redirect(url_for("children_list"))
    classes = class_service.list_classes()
    if request.method == "POST":
        child_service.update_child(
            child_id, request.form["name"], request.form["gender"], request.form["birthday"],
            request.form.get("class_id") or None, request.form["guardian"],
            request.form["phone"], request.form.get("allergy", ""),
        )
        flash("修改成功")
        return redirect(url_for("children_list"))
    return render_template("children/form.html", child=child, classes=classes)

@app.route("/children/<int:child_id>/delete", methods=["POST"])
@login_required
@role_required("admin")
def children_delete(child_id):
    child_service.delete_child(child_id)
    flash("已删除")
    return redirect(url_for("children_list"))

@app.route("/children/export")
@login_required
@role_required("admin", "teacher")
def children_export():
    children = child_service.list_children()
    wb = Workbook()
    ws = wb.active
    ws.title = "儿童名单"
    headers = ["ID", "姓名", "性别", "生日", "班级", "监护人", "电话", "过敏史"]
    ws.append(headers)
    for c in children:
        ws.append([c["id"], c["name"], c["gender"], c["birthday"], c["class_name"] or "-", c["guardian"], c["phone"], c["allergy"] or "-"])
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['E'].width = 20
    ws.column_dimensions['F'].width = 15
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="儿童名单.xlsx")

@app.route("/children/import", methods=["POST"])
@login_required
@role_required("admin", "teacher")
def children_import():
    file = request.files.get("file")
    if not file or file.filename == '':
        flash("请选择要上传的 Excel 文件")
        return redirect(url_for("children_list"))
    try:
        wb = load_workbook(file)
        ws = wb.active
        classes = class_service.list_classes()
        class_map = {c["name"]: c["id"] for c in classes}
        success_count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]: continue
            name = str(row[0]).strip()
            gender = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            birthday = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            class_name = str(row[3]).strip() if len(row) > 3 and row[3] else ""
            guardian = str(row[4]).strip() if len(row) > 4 and row[4] else ""
            phone = str(row[5]).strip() if len(row) > 5 and row[5] else ""
            allergy = str(row[6]).strip() if len(row) > 6 and row[6] else ""
            class_id = class_map.get(class_name)
            child_service.add_child(name, gender, birthday, class_id, guardian, phone, allergy)
            success_count += 1
        flash(f"成功导入 {success_count} 名儿童数据！")
    except Exception as e:
        flash(f"导入失败，请检查表格格式：{str(e)}")
    return redirect(url_for("children_list"))

# ---------- 班级管理 ----------
@app.route("/classes", methods=["GET", "POST"])
@login_required
@role_required("admin")
def classes_list():
    if request.method == "POST":
        class_service.add_class(request.form["name"], request.form.get("grade", ""), request.form.get("teacher_name", ""))
        flash("班级已添加")
        return redirect(url_for("classes_list"))
    classes = class_service.list_classes()
    return render_template("classes/list.html", classes=classes)

@app.route("/classes/<int:class_id>/delete", methods=["POST"])
@login_required
@role_required("admin")
def classes_delete(class_id):
    class_service.delete_class(class_id)
    flash("班级已删除")
    return redirect(url_for("classes_list"))

# ---------- 用户管理 ----------
@app.route("/users", methods=["GET", "POST"])
@login_required
@role_required("admin")
def users_list():
    if request.method == "POST":
        username = request.form["username"].strip()
        password = request.form["password"]
        role = request.form.get("role", "teacher")
        real_name = request.form.get("real_name", "")
        child_id = request.form.get("child_id") or None
        if username and password:
            try:
                conn = get_db()
                conn.execute("INSERT INTO users (username, password_hash, role, real_name, child_id) VALUES (?, ?, ?, ?, ?)",
                             (username, generate_password_hash(password), role, real_name, child_id))
                conn.commit()
                conn.close()
                flash("用户添加成功")
            except Exception:
                flash("添加失败，用户名可能已存在")
        return redirect(url_for("users_list"))
    users = user_service.list_users()
    children = child_service.list_children()
    return render_template("users/list.html", users=users, children=children)

@app.route("/users/<int:user_id>/delete", methods=["POST"])
@login_required
@role_required("admin")
def users_delete(user_id):
    if user_id == current_user.id:
        flash("不能删除当前登录的账号！")
    else:
        user_service.delete_user(user_id)
        flash("用户已删除")
    return redirect(url_for("users_list"))

# ---------- 考勤/接送记录 ----------
@app.route("/records")
@login_required
@role_required("admin", "teacher")
def records_list():
    # 默认查询今天（动态时区）
    date_str = request.args.get("date", datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d"))
    records = record_service.list_records(date_str)
    return render_template("records/list.html", records=records, date_str=date_str)

@app.route("/records/add", methods=["GET", "POST"])
@login_required
@role_required("admin", "teacher")
def records_add():
    if request.method == "POST":
        child_id = request.form.get("child_id")
        record_type = request.form.get("record_type")
        operator = request.form.get("operator")
        if child_id and record_type:
            record_service.add_record(child_id, record_type, operator)
            flash("考勤记录已保存")
            return redirect(url_for("records_list"))
        flash("请选择儿童和记录类型")
    children = child_service.list_children()
    return render_template("records/form.html", children=children)

@app.route("/records/<int:record_id>/delete", methods=["POST"])
@login_required
@role_required("admin")
def records_delete(record_id):
    record_service.delete_record(record_id)
    flash("记录已删除")
    return redirect(url_for("records_list"))

@app.route("/my-records")
@login_required
@role_required("parent")
def my_records():
    if not current_user.child_id:
        flash("您的账号还没有绑定孩子，请联系老师或管理员。")
        return redirect(url_for("index"))
    date_str = request.args.get("date")
    conn = get_db()
    sql = "SELECT * FROM records WHERE child_id = ?"
    params = [current_user.child_id]
    if date_str:
        sql += " AND record_time LIKE ?"
        params.append(f"{date_str}%")
    sql += " ORDER BY record_time DESC"
    records = conn.execute(sql, params).fetchall()
    conn.close()
    return render_template("my_records.html", records=records, date_str=date_str)

# ---------- 通知公告 ----------
def get_notices():
    conn = get_db()
    rows = conn.execute("SELECT * FROM notices ORDER BY id DESC LIMIT 5").fetchall()
    conn.close()
    return rows

@app.route("/notices", methods=["GET", "POST"])
@login_required
@role_required("admin")
def notices_list():
    if request.method == "POST":
        title = request.form["title"]
        content = request.form["content"]
        # 动态时区
        now_str = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")
        conn = get_db()
        conn.execute("INSERT INTO notices (title, content, created_at, author) VALUES (?, ?, ?, ?)",
                     (title, content, now_str, current_user.username))
        conn.commit()
        conn.close()
        flash("通知已发布")
        return redirect(url_for("notices_list"))
    conn = get_db()
    notices = conn.execute("SELECT * FROM notices ORDER BY id DESC").fetchall()
    conn.close()
    return render_template("notices/list.html", notices=notices)

@app.route("/notices/<int:notice_id>/delete", methods=["POST"])
@login_required
@role_required("admin")
def notices_delete(notice_id):
    conn = get_db()
    conn.execute("DELETE FROM notices WHERE id = ?", (notice_id,))
    conn.commit()
    conn.close()
    flash("通知已删除")
    return redirect(url_for("notices_list"))

# ---------- 照片墙/成长记录 ----------
@app.route("/photos", methods=["GET", "POST"])
@login_required
@role_required("admin", "teacher")
def photos_list():
    if request.method == "POST":
        child_id = request.form.get("child_id")
        description = request.form.get("description", "")
        file = request.files.get("image")
        
        if not child_id or not file or file.filename == '':
            flash("请选择儿童并上传照片")
            return redirect(url_for("photos_list"))
            
        ext = os.path.splitext(file.filename)[1]
        filename = f"{uuid.uuid4().hex}{ext}"
        save_path = os.path.join(app.config['UPLOAD_FOLDER'], filename)
        file.save(save_path)
        
        db_path = f"uploads/{filename}"
        # 动态时区
        now_str = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")
        
        conn = get_db()
        conn.execute("INSERT INTO photos (child_id, image_path, description, created_at, author) VALUES (?, ?, ?, ?, ?)",
                     (child_id, db_path, description, now_str, current_user.username))
        conn.commit()
        conn.close()
        
        flash("照片上传成功")
        return redirect(url_for("photos_list"))
        
    children = child_service.list_children()
    photos = get_all_photos()
    return render_template("photos/list.html", children=children, photos=photos)

@app.route("/photos/<int:photo_id>/delete", methods=["POST"])
@login_required
@role_required("admin", "teacher")
def photos_delete(photo_id):
    conn = get_db()
    row = conn.execute("SELECT image_path FROM photos WHERE id = ?", (photo_id,)).fetchone()
    if row:
        try:
            file_path = os.path.join(app.config['UPLOAD_FOLDER'], os.path.basename(row["image_path"]))
            if os.path.exists(file_path):
                os.remove(file_path)
        except Exception:
            pass
        conn.execute("DELETE FROM photos WHERE id = ?", (photo_id,))
        conn.commit()
    conn.close()
    flash("照片已删除")
    return redirect(url_for("photos_list"))

def get_all_photos():
    conn = get_db()
    rows = conn.execute("""SELECT p.*, c.name AS child_name 
                           FROM photos p JOIN children c ON p.child_id = c.id 
                           ORDER BY p.id DESC LIMIT 20""").fetchall()
    conn.close()
    return rows

def get_photos_by_child(child_id):
    conn = get_db()
    rows = conn.execute("SELECT * FROM photos WHERE child_id = ? ORDER BY id DESC LIMIT 10", (child_id,)).fetchall()
    conn.close()
    return rows

# ---------- 请假管理 ----------
@app.route("/leaves", methods=["GET", "POST"])
@login_required
@role_required("admin", "teacher")
def leaves_list():
    if request.method == "POST":
        leave_id = request.form.get("leave_id")
        status = request.form.get("status")
        if leave_id and status in ("已同意", "已驳回"):
            conn = get_db()
            conn.execute("UPDATE leaves SET status = ? WHERE id = ?", (status, leave_id))
            conn.commit()
            conn.close()
            flash("审批状态已更新")
        return redirect(url_for("leaves_list"))
        
    conn = get_db()
    leaves = conn.execute("""SELECT l.*, c.name AS child_name, cl.name AS class_name
                             FROM leaves l JOIN children c ON l.child_id = c.id
                             LEFT JOIN classes cl ON c.class_id = cl.id
                             ORDER BY l.id DESC""").fetchall()
    conn.close()
    return render_template("leaves/list.html", leaves=leaves)

@app.route("/my-leaves", methods=["GET", "POST"])
@login_required
@role_required("parent")
def my_leaves():
    if not current_user.child_id:
        flash("您的账号还没有绑定孩子，无法请假。")
        return redirect(url_for("index"))
        
    if request.method == "POST":
        start_date = request.form["start_date"]
        end_date = request.form["end_date"]
        reason = request.form["reason"]
        # 动态时区
        now_str = datetime.now(ZoneInfo("Asia/Shanghai")).strftime("%Y-%m-%d %H:%M")
        
        conn = get_db()
        conn.execute("INSERT INTO leaves (child_id, start_date, end_date, reason, status, created_at) VALUES (?, ?, ?, ?, '待审批', ?)",
                     (current_user.child_id, start_date, end_date, reason, now_str))
        conn.commit()
        conn.close()
        flash("请假申请已提交，请等待老师审批")
        return redirect(url_for("my_leaves"))
        
    conn = get_db()
    leaves = conn.execute("SELECT * FROM leaves WHERE child_id = ? ORDER BY id DESC", (current_user.child_id,)).fetchall()
    conn.close()
    return render_template("my_leaves.html", leaves=leaves)


if __name__ == "__main__":
    init_db()
    app.run(debug=True)

if __name__ != "__main__":
    init_db()